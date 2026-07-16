#!/usr/bin/env python3
"""
stream_xref_disk_gen_v3.py
==========================
Generate RAW forensic benchmark disk images — NO filesystem whatsoever.

IMAGE LAYOUT
  Grid          : 24 images  (6 scenarios × 2 sizes × 2 frag-counts)
  Total         : 24 images

NO boot sector · NO FAT · NO directory entries · NO partition table
Image = pseudo-random noise + PDF bytes written at sector-aligned offsets.
Identical raw-dd format to DFRWS 2007 challenge image.

DISK GEOMETRY  (matches real HDD sector geometry)
  Sector size   : 512 bytes   (industry standard, same as DFRWS)
  Cluster/block : 4096 bytes  (8 sectors, typical HDD allocation unit)
  No filesystem structures — only sector offsets matter to carvers

Seed : 42 (per image: 42 + scenario_index × 1000 + size_mb + n_frags)
"""

import argparse, math, mmap, os, random, struct, sys
from pathlib import Path
from typing import List, Tuple

# Force UTF-8 stdout so the box-drawing/arrow characters in the progress output
# below don't crash on a plain Windows console (legacy cp1252 codepage).
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# ── pip install openpyxl  (standard; ships with Anaconda) ─────────────────────
try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
except ImportError:
    sys.exit("ERROR: openpyxl not found. Run:  pip install openpyxl")

# ── Paths (set from --src/--out in main(); no personal-machine defaults) ──────
DIR_PDF  = None
DIR_OUT  = None
XLSX_OUT = None

# ── Grid ──────────────────────────────────────────────────────────────────────
SIZES_MB    = [64, 256]
FRAG_COUNTS = [3, 8]
SCENARIOS   = ["CONT", "SEQF", "NSQF", "MISS", "NEST", "BRAID"]
# → 6 × 2 × 2 = 24 images

# ── Disk geometry (matches real HDD — no filesystem) ─────────────────────────
SEC_SZ  = 512           # bytes per sector  (HDD standard)
BLK_SZ  = 4096          # bytes per block   (= 8 sectors, typical HDD cluster size)
SPB     = BLK_SZ // SEC_SZ   # sectors per block = 8
# No filesystem: no MBR, no partition table, no FAT, no directory entries
# Carvers see: random noise background + PDF bytes at known sector offsets

MAX_EXCEL_FRAGS = 8     # frag sector-range columns in Excel: Frag1..Frag8

# ─────────────────────────────────────────────────────────────────────────────
# Raw disk builder — NO filesystem, NO FAT, NO boot sector
# Identical format to DFRWS 2007 (raw dd image).
# Each write_* method returns:
#   (first_block, List[Tuple[start_sector, end_sector]])  per logical fragment
# ─────────────────────────────────────────────────────────────────────────────
class RawDisk:

    def __init__(self, path: Path, size_mb: int):
        self.path    = path
        self.size    = size_mb * 1024 * 1024
        self.tot_sec = self.size // SEC_SZ
        self.tot_blk = self.size // BLK_SZ

        # Fill entire image with pseudo-random noise using the seeded `random`
        # module (deterministic given main()'s random.seed() call, unlike
        # os.urandom which is OS entropy and can never be reproduced run-to-run).
        # Carvers must find %PDF- headers among noise — realistic background.
        print(f"  [{size_mb:>4d} MB] Writing noise ({size_mb} MB) …", end="", flush=True)
        CHUNK = 4 * 1024 * 1024   # 4 MB chunks — fast I/O
        with open(path, 'wb') as f:
            written = 0
            while written < self.size:
                n = min(CHUNK, self.size - written)
                f.write(random.randbytes(n))
                written += n
        print(" done")

        self._f  = open(path, 'r+b')
        self._mm = mmap.mmap(self._f.fileno(), 0)
        self._ptr = 0   # next free block index
        print(f"  [{size_mb:>4d} MB] blocks={self.tot_blk}  sec={self.tot_sec}  (NO filesystem)")

    # ── Geometry helpers ──────────────────────────────────────────────────────
    def _blk_to_sec(self, b: int) -> int:
        """First sector of block b."""
        return b * SPB

    def _blk_range_secs(self, blk_list: List[int]) -> Tuple[int, int]:
        """(start_sector, end_sector_inclusive) for a contiguous block run."""
        s = self._blk_to_sec(blk_list[0])
        e = self._blk_to_sec(blk_list[-1]) + SPB - 1
        return s, e

    def _blk_offset(self, b: int) -> int:
        return b * BLK_SZ

    # ── Low-level I/O ─────────────────────────────────────────────────────────
    def _write_blk(self, b: int, data: bytes):
        off = self._blk_offset(b)
        if off + BLK_SZ > self.size: return
        chunk = (data + bytes(BLK_SZ))[:BLK_SZ]
        self._mm[off: off + BLK_SZ] = chunk

    def _alloc(self) -> int:
        if self._ptr >= self.tot_blk:
            raise RuntimeError("Disk full")
        b = self._ptr; self._ptr += 1; return b

    # ── Scenario 1 : CONT – Non-fragmented ───────────────────────────────────
    def write_contiguous(self, data: bytes) -> Tuple[int, List[Tuple[int,int]]]:
        n = max(1, math.ceil(len(data) / BLK_SZ))
        blk_list = []
        for i in range(n):
            b = self._alloc()
            self._write_blk(b, data[i * BLK_SZ:(i+1) * BLK_SZ])
            blk_list.append(b)
        return blk_list[0], [self._blk_range_secs(blk_list)]

    # ── Scenario 2 : SEQF – Sequential fragments ──────────────────────────────
    def write_seq_frag(self, data: bytes, n_frags: int = 3,
                       gap: int = 5) -> Tuple[int, List[Tuple[int,int]]]:
        n     = max(1, math.ceil(len(data) / BLK_SZ))
        parts = _splitN(n, n_frags)
        frag_secs: List[Tuple[int,int]] = []
        off = 0; first = None

        for part_sz in parts:
            frag_blks = []
            for i in range(part_sz):
                b = self._alloc()
                self._write_blk(b, data[off * BLK_SZ:(off+1) * BLK_SZ])
                frag_blks.append(b); off += 1
            if first is None: first = frag_blks[0]
            frag_secs.append(self._blk_range_secs(frag_blks))
            # gap blocks: overwrite with fresh random noise (other-file simulation)
            for _ in range(gap):
                b = self._alloc()
                self._write_blk(b, random.randbytes(BLK_SZ))

        return first, frag_secs

    # ── Scenario 3 : NSQF – Non-sequential fragments ──────────────────────────
    def write_nonseq_frag(self, data: bytes,
                          n_frags: int = 3) -> Tuple[int, List[Tuple[int,int]]]:
        n     = max(1, math.ceil(len(data) / BLK_SZ))
        parts = _splitN(n, n_frags)

        # Allocate all blocks upfront (contiguous on disk)
        blocks = [self._alloc() for _ in range(n)]

        # Split into fragment groups
        frags: List[List[int]] = []
        idx = 0
        for sz in parts:
            frags.append(blocks[idx: idx+sz]); idx += sz

        # Shuffle fragment order → non-sequential layout
        frag_order = list(range(len(frags)))
        if len(frag_order) > 1:
            for _ in range(20):
                random.shuffle(frag_order)
                if frag_order != sorted(frag_order): break

        # Write data into shuffled block positions
        data_off = 0; first = None
        frag_secs: List[Tuple[int,int]] = []
        for fi in frag_order:
            frag = frags[fi]
            for b in frag:
                self._write_blk(b, data[data_off * BLK_SZ:(data_off+1) * BLK_SZ])
                data_off += 1
            if first is None: first = frag[0]
            frag_secs.append(self._blk_range_secs(frag))

        return first, frag_secs

    # ── Scenario 4 : MISS – Missing fragment ──────────────────────────────────
    def write_missing_frag(self, data: bytes) -> Tuple[int, List[Tuple[int,int]]]:
        n = max(1, math.ceil(len(data) / BLK_SZ))
        blk_list = []
        for i in range(n):
            b = self._alloc()
            self._write_blk(b, data[i * BLK_SZ:(i+1) * BLK_SZ])
            blk_list.append(b)

        # Overwrite ~2/3 tail blocks with noise (missing fragment = overwritten)
        steps   = max(1, n // 3)
        for b in blk_list[steps:]:
            self._write_blk(b, random.randbytes(BLK_SZ))   # destroy tail data

        surviving = blk_list[:steps]
        return blk_list[0], [self._blk_range_secs(surviving)]

    # ── Scenario 5 : NEST – Nested document ───────────────────────────────────
    def write_nested(self, outer: bytes,
                     inner: bytes) -> Tuple[int, List[Tuple[int,int]]]:
        pos     = len(outer) // 3
        payload = outer[:pos] + inner + outer[pos:]
        return self.write_contiguous(payload)

    # ── Scenario 6 : BRAID – Braided documents ────────────────────────────────
    def write_braided(self, a: bytes, b: bytes
                      ) -> Tuple[Tuple[int, List[Tuple[int,int]]],
                                 Tuple[int, List[Tuple[int,int]]]]:
        na = max(1, math.ceil(len(a) / BLK_SZ))
        nb = max(1, math.ceil(len(b) / BLK_SZ))
        ba_list: List[int] = []
        bb_list: List[int] = []

        for i in range(max(na, nb)):
            if i < na:
                blk = self._alloc()
                self._write_blk(blk, a[i * BLK_SZ:(i+1) * BLK_SZ])
                ba_list.append(blk)
            if i < nb:
                blk = self._alloc()
                self._write_blk(blk, b[i * BLK_SZ:(i+1) * BLK_SZ])
                bb_list.append(blk)

        def single_blk_secs(lst):
            return [(self._blk_to_sec(b), self._blk_to_sec(b) + SPB - 1) for b in lst]

        return (ba_list[0], single_blk_secs(ba_list)), \
               (bb_list[0], single_blk_secs(bb_list))

    # ── Finalize ──────────────────────────────────────────────────────────────
    def finalize(self):
        """No filesystem to write — just flush and close."""
        self._mm.flush()
        self._mm.close()
        self._f.close()
        sz = self.path.stat().st_size
        print(f"  ✓ {self.path.name}  ({sz:,} B  blocks_used={self._ptr}/{self.tot_blk}  NO filesystem)")

    def __del__(self):
        try:
            if not self._mm.closed: self._mm.close()
            if not self._f.closed:  self._f.close()
        except Exception: pass


# ── Helpers ───────────────────────────────────────────────────────────────────
def _splitN(n: int, parts: int) -> List[int]:
    parts = max(1, min(parts, n))
    base, rem = divmod(n, parts)
    return [base + (1 if i < rem else 0) for i in range(parts)]


# ── PDF loader ────────────────────────────────────────────────────────────────
def load_pdfs() -> List[Tuple[str, bytes]]:
    """
    Load exactly 20 PDFs from DIR_PDF in sorted order.
    Expected composition (per previous conversation):
      ~6  classic xref-table PDFs       (PDF 1.2–1.4)
      ~6  xref-stream PDFs              (PDF 1.5+)     → C6 target
      ~4  PNG-predictor + xref-stream   (hardest C6)
      ~2–3 linearized / hybrid-ref
      ~1–2 embedded files (font+image)  → C5 render stress
    All PDFs >= 500 KB for realistic fragmentation.
    """
    pdfs = []
    for p in sorted(DIR_PDF.glob("*.pdf")):
        data = p.read_bytes()
        pdfs.append((p.name, data))
        print(f"  PDF: {p.name}  ({len(data):,} B)")
    if len(pdfs) < 20:
        sys.exit(f"ERROR: need 20 PDFs in {DIR_PDF}, found {len(pdfs)}")
    if len(pdfs) > 20:
        print(f"WARNING: {len(pdfs)} PDFs found — using first 20 (sorted)")
        pdfs = pdfs[:20]
    # Warn on small PDFs
    small = [(nm, sz) for nm, d in pdfs for sz in [len(d)] if sz < 512*1024]
    if small:
        for nm, sz in small:
            print(f"  WARNING: {nm} is {sz//1024} KB — may produce shallow fragmentation")
    print(f"Loaded: {len(pdfs)} PDFs  "
          f"(total {sum(len(d) for _,d in pdfs)/1024/1024:.1f} MB)")
    return pdfs


# ── Single image builder ──────────────────────────────────────────────────────
def build_image(image_id: str, scenario: str, size_mb: int,
                n_frags: int, pdfs: List[Tuple[str,bytes]]
                ) -> Tuple[Path, List[dict]]:
    """
    Returns (image_path, list_of_manifest_rows).
    Each manifest_row dict has keys:
      image_id, scenario, size_mb, n_frags, pdf_name, pdf_size,
      first_sector, last_sector, frag_sectors  (list of (s,e) tuples)
    """
    size_label = f"{size_mb}MB"
    fname  = f"{image_id}.img"
    fpath  = DIR_OUT / fname
    disk   = RawDisk(fpath, size_mb)
    rows: List[dict] = []

    def add_row(pdf_name, pdf_size, first_clus, frag_secs):
        # first_sector = start of first fragment; last_sector = end of last fragment
        all_s = [s for s,e in frag_secs]
        all_e = [e for s,e in frag_secs]
        rows.append({
            "image_id"    : image_id,
            "scenario"    : scenario,
            "size_label"  : size_label,
            "n_frags"     : n_frags,
            "pdf_name"    : pdf_name,
            "pdf_size"    : pdf_size,
            "first_sector": min(all_s),
            "last_sector" : max(all_e),
            "frag_sectors": frag_secs,   # full list; Excel shows up to 8
        })

    # ── CONT ─────────────────────────────────────────────────────────────────
    if scenario == "CONT":
        for k, (nm, data) in enumerate(pdfs):
            try:
                fc, fs = disk.write_contiguous(data)
                add_row(nm, len(data), fc, fs)
                print(f"    [CONT {k:02d}] sec={fs[0][0]}–{fs[0][1]}  {nm}")
            except RuntimeError as e: print(f"    [CONT {k:02d}] SKIP {e}")

    # ── SEQF ─────────────────────────────────────────────────────────────────
    elif scenario == "SEQF":
        for k, (nm, data) in enumerate(pdfs):
            try:
                fc, fs = disk.write_seq_frag(data, n_frags=n_frags, gap=5)
                add_row(nm, len(data), fc, fs)
                print(f"    [SEQF {k:02d}] frags={len(fs)}  {nm}")
            except RuntimeError as e: print(f"    [SEQF {k:02d}] SKIP {e}")

    # ── NSQF ─────────────────────────────────────────────────────────────────
    elif scenario == "NSQF":
        for k, (nm, data) in enumerate(pdfs):
            try:
                fc, fs = disk.write_nonseq_frag(data, n_frags=n_frags)
                add_row(nm, len(data), fc, fs)
                print(f"    [NSQF {k:02d}] frags={len(fs)}  {nm}")
            except RuntimeError as e: print(f"    [NSQF {k:02d}] SKIP {e}")

    # ── MISS ─────────────────────────────────────────────────────────────────
    elif scenario == "MISS":
        for k, (nm, data) in enumerate(pdfs):
            try:
                fc, fs = disk.write_missing_frag(data)
                add_row(nm, len(data), fc, fs)
                print(f"    [MISS {k:02d}] surviving sec={fs[0][0]}–{fs[0][1]}  {nm}")
            except RuntimeError as e: print(f"    [MISS {k:02d}] SKIP {e}")

    # ── NEST (pairs: 0+1, 2+3, …) ────────────────────────────────────────────
    # NEST embeds inner PDF inside outer. Both are GT targets.
    # Emit TWO manifest rows per pair so GT count = 20, not 10.
    elif scenario == "NEST":
        k = 0
        while k < len(pdfs):
            if k + 1 < len(pdfs):
                outer_nm, outer_d = pdfs[k]
                inner_nm, inner_d = pdfs[k+1]
                try:
                    fc, fs = disk.write_nested(outer_d, inner_d)
                    # Row 1: outer PDF — starts at fragment start, full outer size
                    add_row(outer_nm, len(outer_d), fc, fs)
                    # Row 2: inner PDF — same sector range (embedded inside outer)
                    # Size = inner only; sector range same (both inside same blocks)
                    add_row(inner_nm, len(inner_d), fc, fs)
                    print(f"    [NEST {k//2:02d}] outer={outer_nm} inner={inner_nm}  (2 GT rows)")
                except RuntimeError as e: print(f"    [NEST {k//2:02d}] SKIP {e}")
                k += 2
            else:
                nm, data = pdfs[k]
                try:
                    fc, fs = disk.write_contiguous(data)
                    add_row(nm, len(data), fc, fs)
                except RuntimeError: pass
                k += 1

    # ── BRAID (pairs: 0+1, 2+3, …) ───────────────────────────────────────────
    elif scenario == "BRAID":
        k = 0
        while k < len(pdfs):
            if k + 1 < len(pdfs):
                nm_a, data_a = pdfs[k]
                nm_b, data_b = pdfs[k+1]
                try:
                    (fa, fs_a), (fb, fs_b) = disk.write_braided(data_a, data_b)
                    add_row(nm_a, len(data_a), fa, fs_a)
                    add_row(nm_b, len(data_b), fb, fs_b)
                    print(f"    [BRAID {k//2:02d}] {nm_a} ↔ {nm_b}  clusters_interleaved")
                except RuntimeError as e: print(f"    [BRAID {k//2:02d}] SKIP {e}")
                k += 2
            else:
                nm, data = pdfs[k]
                try:
                    fc, fs = disk.write_contiguous(data)
                    add_row(nm, len(data), fc, fs)
                except RuntimeError: pass
                k += 1

    disk.finalize()

    # Per-image CSV manifest (full sector detail)
    csv_path = DIR_OUT / f"{image_id}_manifest.csv"
    with open(csv_path, 'w', newline='') as f:
        # build header
        frag_hdrs = []
        for i in range(1, MAX_EXCEL_FRAGS+1):
            frag_hdrs += [f"frag{i}_start_sec", f"frag{i}_end_sec"]
        f.write("image_id,scenario,size,n_frags,pdf_name,pdf_size_bytes,"
                "first_sector,last_sector," + ",".join(frag_hdrs) + "\n")
        for r in rows:
            fs_cells = []
            for i in range(MAX_EXCEL_FRAGS):
                if i < len(r["frag_sectors"]):
                    s, e = r["frag_sectors"][i]
                    fs_cells += [str(s), str(e)]
                else:
                    fs_cells += ["", ""]
            f.write(",".join([
                r["image_id"], r["scenario"], r["size_label"], str(r["n_frags"]),
                r["pdf_name"], str(r["pdf_size"]),
                str(r["first_sector"]), str(r["last_sector"])
            ] + fs_cells) + "\n")
    print(f"  CSV: {csv_path.name}")
    return fpath, rows


# ── Excel builder ─────────────────────────────────────────────────────────────
def write_excel(all_rows: List[dict]):
    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary (one row per image) ──────────────────────────────────
    ws_sum = wb.active; ws_sum.title = "Image Summary"
    sum_headers = ["Image ID", "Scenario", "Size", "Frags", "PDFs", "Image File"]
    _write_header_row(ws_sum, 1, sum_headers)

    # group rows by image_id
    from collections import defaultdict, OrderedDict
    by_image: dict = OrderedDict()
    for r in all_rows:
        by_image.setdefault(r["image_id"], []).append(r)

    for ri, (iid, rows) in enumerate(by_image.items(), start=2):
        r0 = rows[0]
        _write_data_row(ws_sum, ri, [
            iid, r0["scenario"], r0["size_label"], r0["n_frags"],
            len(rows), f"{iid}.img"
        ], alt=(ri % 2 == 0))

    for col in ws_sum.columns:
        ws_sum.column_dimensions[col[0].column_letter].width = 22
    ws_sum.freeze_panes = "A2"

    # ── Sheet 2: Full manifest (one row per PDF per image) ────────────────────
    ws = wb.create_sheet("Sector Manifest")
    frag_hdrs = []
    for i in range(1, MAX_EXCEL_FRAGS+1):
        frag_hdrs += [f"Frag{i} Start Sec", f"Frag{i} End Sec"]
    headers = (["Image ID", "Scenario", "Size", "Frags", "PDF Name",
                "PDF Size (B)", "First Sector", "Last Sector"] + frag_hdrs)
    _write_header_row(ws, 1, headers)

    ri = 2
    for r in all_rows:
        frag_cells = []
        for i in range(MAX_EXCEL_FRAGS):
            if i < len(r["frag_sectors"]):
                s, e = r["frag_sectors"][i]
                frag_cells += [s, e]
            else:
                frag_cells += ["", ""]
        _write_data_row(ws, ri, [
            r["image_id"], r["scenario"], r["size_label"], r["n_frags"],
            r["pdf_name"], r["pdf_size"],
            r["first_sector"], r["last_sector"]
        ] + frag_cells, alt=(ri % 2 == 0))
        ri += 1

    # column widths
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 7
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 14
    for i in range(MAX_EXCEL_FRAGS * 2):
        col_letter = openpyxl.utils.get_column_letter(9 + i)
        ws.column_dimensions[col_letter].width = 15
    ws.freeze_panes = "A2"

    wb.save(XLSX_OUT)
    print(f"\nExcel manifest: {XLSX_OUT}")


# ── Excel styling helpers ─────────────────────────────────────────────────────
HEAD_FILL  = PatternFill("solid", fgColor="2E5A88")
HEAD_FONT  = Font(bold=True, color="FFFFFF", name="Arial", size=10)
ALT_FILL   = PatternFill("solid", fgColor="EAF1F8")
NORM_FILL  = PatternFill("solid", fgColor="FFFFFF")
DATA_FONT  = Font(name="Arial", size=9)
THIN       = Side(style="thin", color="B0BEC5")
BORDER     = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER     = Alignment(horizontal="center", vertical="center")
LEFT       = Alignment(horizontal="left",   vertical="center")

def _write_header_row(ws, row: int, headers: List[str]):
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=ci, value=h)
        c.fill = HEAD_FILL; c.font = HEAD_FONT
        c.alignment = CENTER; c.border = BORDER
    ws.row_dimensions[row].height = 18

def _write_data_row(ws, row: int, values: list, alt: bool = False):
    fill = ALT_FILL if alt else NORM_FILL
    for ci, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=ci, value=v)
        c.fill = fill; c.font = DATA_FONT; c.border = BORDER
        c.alignment = LEFT if ci <= 2 else CENTER


# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    global DIR_PDF, DIR_OUT, XLSX_OUT
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True, help="folder containing the 20 source PDFs")
    ap.add_argument("--out", required=True, help="output folder for disk images")
    args = ap.parse_args()
    DIR_PDF = Path(args.src)
    DIR_OUT = Path(args.out)
    XLSX_OUT = DIR_OUT / "disk_image_manifest.xlsx"

    DIR_OUT.mkdir(parents=True, exist_ok=True)

    pdfs = load_pdfs()

    total_pdf_mb = sum(len(d) for _, d in pdfs) / (1024*1024)
    print(f"\nPDF corpus: {total_pdf_mb:.1f} MB across {len(pdfs)} files")
    print(f"Output dir: {DIR_OUT}\n")

    all_rows: List[dict] = []
    img_num = 0

    for scenario in SCENARIOS:
        scen_idx = SCENARIOS.index(scenario)
        for size_mb in SIZES_MB:
            for n_frags in FRAG_COUNTS:
                img_num += 1
                size_label = f"{size_mb}MB"
                image_id   = f"SYN-{scenario}-{size_label}-F{n_frags}"

                print(f"\n{'═'*64}")
                print(f"[{img_num:02d}/24] Building: {image_id}")
                print(f"  scenario={scenario}  size={size_mb}MB  frags={n_frags}")

                # Estimate fit: fragmented scenarios need ~2× raw PDF bytes
                needed_mb = total_pdf_mb * 2.2
                if needed_mb > size_mb * 0.80:
                    print(f"  WARNING: PDFs ({total_pdf_mb:.1f} MB ×2.2) may be tight for {size_mb} MB")

                # Per-image seed: deterministic but varies by scenario+size+frags
                random.seed(42 + scen_idx * 1000 + size_mb + n_frags)

                try:
                    _, rows = build_image(image_id, scenario, size_mb, n_frags, pdfs)
                    all_rows.extend(rows)
                except Exception as exc:
                    import traceback
                    print(f"  ERROR: {exc}")
                    traceback.print_exc()

    print(f"\n{'═'*64}")
    print(f"All 24 images built. Writing Excel manifest…")
    write_excel(all_rows)

    print(f"\n{'═'*64}")
    print("Done.")
    print(f"  Images : {DIR_OUT}")
    print(f"  Excel  : {XLSX_OUT}")
    print(f"  Scoring rows in manifest: {len(all_rows)}  (24 images × 20 PDFs)")


if __name__ == "__main__":
    main()
